#!/usr/bin/env python3
"""Generate the ZeoFex × John Paul Kpantey VentraPOS partner agreement as a Word file."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(
    "/Users/insight/Documents/ventrav1/ZeoFex-VentraPOS-Partner-Agreement-John-Paul-Kpantey.docx"
)

GREEN = RGBColor(0x00, 0x3D, 0x2E)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x44, 0x44, 0x44)
HEADER_FILL = "003D2E"
ROW_FILL = "F3F7F5"


def set_run_font(run, *, name="Times New Roman", size=12, bold=False, italic=False, color=BLACK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag == qn("w:shd"):
            tc_pr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_borders(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "003D2E")
        borders.append(el)
    tc_pr.append(borders)


def set_cell_text(cell, text, *, bold=False, color=BLACK, size=10, align="left"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    set_cell_borders(cell)


def add_para(doc, text, *, size=12, bold=False, italic=False, align="justify", space_after=8, space_before=0, color=BLACK, first_line=None):
    p = doc.add_paragraph()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading_styled(doc, text, *, size=13):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper())
    set_run_font(run, size=size, bold=True, color=GREEN)
    # underline via bottom border on paragraph
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "003D2E")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    return p


def add_subhead(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=True, color=GREEN)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run, size=12)
    return p


def make_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, h, bold=True, color=RGBColor(255, 255, 255), size=10, align="left")
        shade_cell(cell, HEADER_FILL)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            set_cell_text(cell, val, size=10)
            if r_i % 2 == 1:
                shade_cell(cell, ROW_FILL)
    doc.add_paragraph()
    return table


def add_blank_line(doc, label, width=52):
    add_para(doc, f"{label} {'_' * width}", align="left", space_after=10, size=12)


def add_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run("ZeoFex Technology  ·  VentraPOS Channel Partner Agreement  ·  Page ")
    set_run_font(run, size=8, color=GRAY)

    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r2 = p.add_run()
    r2._r.append(fld1)
    r2._r.append(instr)
    r2._r.append(fld2)
    set_run_font(r2, size=8, color=GRAY)

    run3 = p.add_run(" of ")
    set_run_font(run3, size=8, color=GRAY)

    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "begin")
    instr2 = OxmlElement("w:instrText")
    instr2.set(qn("xml:space"), "preserve")
    instr2.text = " NUMPAGES "
    fld4 = OxmlElement("w:fldChar")
    fld4.set(qn("w:fldCharType"), "end")
    r4 = p.add_run()
    r4._r.append(fld3)
    r4._r.append(instr2)
    r4._r.append(fld4)
    set_run_font(r4, size=8, color=GRAY)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.2)
    add_page_number(section)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.font.color.rgb = BLACK

    # Cover / title
    add_para(doc, "ZEOFEX TECHNOLOGY", size=11, bold=True, align="center", space_after=2, color=GREEN)
    add_para(doc, "Trading as VentraPOS", size=11, italic=True, align="center", space_after=14, color=GRAY)
    add_para(
        doc,
        "CHANNEL PARTNER REVENUE-SHARE AGREEMENT",
        size=18,
        bold=True,
        align="center",
        space_after=6,
        color=GREEN,
    )
    add_para(
        doc,
        "Lifetime Originating Partner — Shops Brought On Board",
        size=12,
        italic=True,
        align="center",
        space_after=6,
        color=GRAY,
    )
    add_para(
        doc,
        "ZeoFex Technology  and  Mr. John Paul Kpantey",
        size=12,
        bold=True,
        align="center",
        space_after=4,
    )
    add_para(
        doc,
        "Governing law: Republic of Ghana",
        size=10,
        align="center",
        space_after=18,
        color=GRAY,
    )

    add_para(
        doc,
        "This Channel Partner Revenue-Share Agreement (the “Agreement”) is made on ________________________ (the “Effective Date”)",
        space_after=12,
    )
    add_para(doc, "BETWEEN", bold=True, align="center", space_after=10)
    add_para(
        doc,
        "ZeoFex Technology, trading as VentraPOS, of Abuakwa-Tanoso, Despite & Kessben Street, Ghana, email zeofex.tech@gmail.com, phone +233 242 975 483 (the “Company”);",
        space_after=10,
    )
    add_para(doc, "AND", bold=True, align="center", space_after=10)
    add_para(
        doc,
        "Mr. John Paul Kpantey, of ________________________________, Ghana Card / ID No. ____________________, TIN ____________________, email ____________________, phone ____________________ (the “Partner”).",
        space_after=10,
    )
    add_para(
        doc,
        "The Company and the Partner are each a “Party” and together the “Parties”.",
        space_after=8,
    )

    # 1
    add_heading_styled(doc, "1. Background")
    add_para(
        doc,
        "1.1 The Company owns and operates VentraPOS, a cloud point-of-sale and business operations platform for retailers and similar businesses (the “Platform”).",
    )
    add_para(
        doc,
        "1.2 The Partner wishes to source, introduce, and help onboard shops onto the Platform.",
    )
    add_para(
        doc,
        "1.3 For each shop the Partner brings on board, the Partner shall receive fifty percent (50%) of that shop’s Net Distributable Revenue for the life of that shop on VentraPOS, after the deductions in this Agreement, including payment fees, platform resources, and the Maintenance Fee.",
    )
    add_para(
        doc,
        "1.4 All software updates, upgrades, patches, and new Platform features shall be designed, built, released, and owned solely by the Company. The Partner is not required to build, maintain, or update the software. The cost of that ongoing work is recovered through the Maintenance Fee and other Resource Costs in Schedule A, which are deducted before the 50/50 split.",
    )
    add_para(
        doc,
        "1.5 This Agreement does not give the Partner any shareholding, ownership, or control of ZeoFex Technology, the Platform, the VentraPOS brand, or shops the Partner did not originate.",
    )

    # 2
    add_heading_styled(doc, "2. Definitions")
    add_para(
        doc,
        "“Referred Shop” means a business account on the Platform that: (a) was first introduced by the Partner; (b) signed up or converted using the Partner’s agreed tracking method (referral code, unique link, written introduction, or written confirmation by the Company); and (c) is recorded on the Company’s Partner Attribution Register.",
    )
    add_para(
        doc,
        "“Life of the Shop” means the entire period during which that Referred Shop remains a VentraPOS customer and, for paid plans, continues to pay. It is not limited to one year, one plan cycle, or the term of this Agreement.",
    )
    add_para(
        doc,
        "“Gross Collected Revenue” means subscription fees, plan fees, branch add-on fees, and other Platform software fees actually received by the Company in Ghana Cedis (GHS) from a Referred Shop in a Settlement Period. It does not include VAT or other taxes collected for the Ghana Revenue Authority, the shop’s own retail sales, or hardware unless listed in Schedule B.",
    )
    add_para(
        doc,
        "“Maintenance Fee” means the Company’s charge for keeping the Platform running and current for that Referred Shop, including software updates, upgrades, bug fixes, security patches, version releases, routine technical upkeep, and related engineering time. The Maintenance Fee is a Resource Cost. It is deducted before the 50/50 split. The Partner does not perform this work and does not receive a separate fee for it.",
    )
    add_para(
        doc,
        "“Resource Costs” means the Company’s actual, reasonable third-party and allocated operating costs of serving that Referred Shop, including the Maintenance Fee, as set out in Schedule A.",
    )
    add_para(
        doc,
        "“Net Distributable Revenue” means, for each Referred Shop in a Settlement Period: Gross Collected Revenue, minus refunds, chargebacks, failed payments, and promotional discounts, minus payment-processor fees, minus Resource Costs (including the Maintenance Fee), minus taxes, levies, or statutory withholdings the Company must deduct.",
    )
    add_para(
        doc,
        "“Partner Share” means 50% of Net Distributable Revenue of each Referred Shop, for the Life of the Shop.",
    )
    add_para(
        doc,
        "“Company Share” means the remaining 50% of Net Distributable Revenue of each Referred Shop.",
    )
    add_para(doc, "“Settlement Period” means each calendar month.")

    # 3
    add_heading_styled(doc, "3. Appointment")
    add_para(
        doc,
        "3.1 The Company appoints Mr. John Paul Kpantey as a non-exclusive originating partner to introduce shops to VentraPOS.",
    )
    add_para(
        doc,
        "3.2 The Partner is an independent contractor, not an employee, director, or person with power to bind the Company, except as written here.",
    )
    add_para(
        doc,
        "3.3 The Partner has no authority to set prices, grant discounts above any written limit, promise features, sign for the Company, collect shop money into a personal account, or release software, unless the Company agrees in writing.",
    )

    # 4
    add_heading_styled(doc, "4. Software updates and maintenance (Company only)")
    add_para(doc, "4.1 The Company shall, at its own operational control:")
    add_bullet(doc, "develop, test, and release all Platform updates and upgrades;")
    add_bullet(doc, "fix bugs and apply security patches;")
    add_bullet(doc, "maintain hosting, database, backups, and related infrastructure; and")
    add_bullet(doc, "decide the product roadmap.")
    add_para(
        doc,
        "4.2 The Partner shall not be asked to code, deploy, or maintain the Platform, and shall not claim ownership of any update.",
    )
    add_para(
        doc,
        "4.3 The Maintenance Fee for each Referred Shop is part of Resource Costs (Schedule A) and is deducted before calculating the Partner Share.",
    )
    add_para(
        doc,
        "4.4 Updates may add, change, or remove features. A Referred Shop’s plan price may change according to the Company’s live price list. The Partner Share remains 50% of Net Distributable Revenue of that shop for the Life of the Shop, including after plan changes.",
    )

    # 5
    add_heading_styled(doc, "5. What the 50% covers (for life)")
    add_para(
        doc,
        "5.1 The 50% applies shop by shop, only to Referred Shops originated by Mr. John Paul Kpantey.",
    )
    add_para(
        doc,
        "5.2 Lifetime rule. For each Referred Shop, the Partner Share continues for the Life of the Shop: for as long as that shop stays on VentraPOS and, where applicable, keeps paying. This includes renewals, upgrades, downgrades, extra paid branches, and later plan changes of that same shop. It does not expire after 12 months or when this Agreement ends, except as in Clauses 5.5 and 15.4.",
    )
    add_para(doc, "5.3 The Partner is not entitled to 50% of:")
    add_bullet(doc, "shops already on the Platform before the Partner introduced them;")
    add_bullet(doc, "shops introduced by the Company or any other person;")
    add_bullet(doc, "Platform-wide revenue, brand value, or company profits; or")
    add_bullet(
        doc,
        "a Starter / free-forever shop, until that shop pays a paid plan or paid add-on.",
    )
    add_para(
        doc,
        "5.4 If a Referred Shop later upgrades from a free plan to a paid plan, lifetime 50% of Net Distributable Revenue starts from the first paid collection and continues for the Life of the Shop.",
    )
    add_para(doc, "5.5 Lifetime attribution stops for a shop if:")
    add_bullet(doc, "the shop permanently closes its VentraPOS account and does not return; or")
    add_bullet(
        doc,
        "the shop churns and later returns without the Partner’s involvement after 90 days off-platform; or",
    )
    add_bullet(
        doc,
        "the Partner is terminated for fraud, data misuse, or the non-solicit breach in Clause 14.",
    )
    add_para(
        doc,
        "5.6 If the Partner dies, Partner Share then due, and future Partner Share on existing Referred Shops, shall be paid to the Partner’s estate or named next of kin on proof reasonably required by the Company, unless the Partner has assigned those rights in writing with the Company’s consent. New shops introduced after death are not attributed.",
    )
    add_para(
        doc,
        "5.7 Current published paid plans (the live VentraPOS price list controls):",
    )
    make_table(
        doc,
        ["Plan", "Monthly", "Annual"],
        [
            ["Starter", "GHS 0 (no share until they pay)", "GHS 0"],
            ["Growth", "GHS 249", "GHS 2,490"],
            ["Pro", "GHS 349", "GHS 3,490"],
            ["Extra branch add-on", "GHS 50 per extra branch per month", "as billed"],
        ],
    )

    # 6
    add_heading_styled(doc, "6. Attribution of shops")
    add_para(
        doc,
        "6.1 A shop is a Referred Shop only if the Company records it as such within 30 days of first paid conversion (or of signup, if the Parties confirm in writing earlier).",
    )
    add_para(
        doc,
        "6.2 If two people claim the same shop, the Company’s written records control. The Company will not double-pay.",
    )
    add_para(
        doc,
        "6.3 Once recorded, attribution is for the Life of the Shop, subject to Clause 5.5.",
    )

    # 7
    add_heading_styled(doc, "7. Partner duties")
    add_para(doc, "Mr. John Paul Kpantey shall:")
    add_bullet(doc, "introduce genuine shops and not invent or recycle accounts;")
    add_bullet(doc, "describe VentraPOS accurately and not mislead on price, features, or uptime;")
    add_bullet(doc, "assist introduction and first contact as reasonably requested;")
    add_bullet(
        doc,
        "not collect or store shop passwords, POS data, or customer personal data beyond what is needed to introduce the shop;",
    )
    add_bullet(doc, "not offer unauthorised discounts, free months, or side deals;")
    add_bullet(
        doc,
        "not promise software updates, custom builds, or maintenance on behalf of the Company;",
    )
    add_bullet(doc, "comply with Ghanaian law, anti-bribery rules, and data-protection duties; and")
    add_bullet(
        doc,
        "not use the VentraPOS or ZeoFex name in a way that suggests the Partner owns the Company.",
    )

    # 8
    add_heading_styled(doc, "8. Company duties")
    add_para(doc, "ZeoFex Technology shall:")
    add_bullet(doc, "provide, update, and maintain the Platform;")
    add_bullet(doc, "bill shops and provide standard merchant support;")
    add_bullet(doc, "keep a Partner Attribution Register;")
    add_bullet(
        doc,
        "issue a monthly settlement statement per Referred Shop, showing Gross Collected Revenue, each deduction (including Maintenance Fee), Net Distributable Revenue, and Partner Share;",
    )
    add_bullet(doc, "pay the Partner Share in accordance with Clause 10; and")
    add_bullet(doc, "remain the contracting party with each shop.")

    # 9
    add_heading_styled(doc, "9. Deductions (resources and maintenance the Company pays)")
    add_para(
        doc,
        "The 50% is calculated after deductions, not on the headline plan price.",
        italic=True,
    )
    add_para(
        doc,
        "Before the 50/50 split, the Company shall deduct from that shop’s Gross Collected Revenue the items below. Deductions are shop-specific where they can be traced, and otherwise allocated on a fair per-shop basis for that Settlement Period.",
    )

    add_subhead(doc, "9.1 Refunds and failed money")
    add_bullet(doc, "full or partial refunds to the shop;")
    add_bullet(doc, "Paystack or bank chargebacks and reversal fees; and")
    add_bullet(doc, "amounts invoiced but not actually received.")

    add_subhead(doc, "9.2 Payment processing")
    add_para(
        doc,
        "Paystack (or successor processor) fees on that shop’s payments, including card, mobile money, and related processor charges actually incurred.",
    )

    add_subhead(doc, "9.3 Platform resource costs (Schedule A)")
    add_para(
        doc,
        "Allocated or actual costs of running that shop on VentraPOS, including:",
    )
    make_table(
        doc,
        ["Resource", "What it covers"],
        [
            [
                "Maintenance Fee",
                "Software updates, upgrades, patches, security fixes, version releases, and routine technical upkeep performed only by the Company",
            ],
            ["Cloud hosting & compute", "Application hosting, bandwidth, serverless execution"],
            ["Database", "PostgreSQL hosting, storage, backups"],
            ["Cache & background jobs", "Redis, queues, scheduled billing jobs"],
            ["Transactional email", "Receipt, OTP, billing, and support emails"],
            ["SMS", "OTP and operational SMS"],
            ["File & media storage", "Logos, product images, uploads"],
            ["Domain, SSL, DNS", "ventrapos.com and related domains"],
            ["Monitoring & error tracking", "Uptime and application monitoring"],
            ["Security & backups", "Backups, restore, essential security tooling"],
        ],
    )

    add_subhead(doc, "9.4 Taxes and statutory amounts")
    add_bullet(doc, "VAT or other tax collected that is not Company income;")
    add_bullet(doc, "withholding tax the Company is legally required to deduct from the Partner Share; and")
    add_bullet(doc, "GRA or other statutory charges attributable to that revenue.")

    add_subhead(doc, "9.5 Discounts and credits")
    add_bullet(doc, "trial periods (no Gross Collected Revenue = no Partner Share);")
    add_bullet(doc, "approved promotional discounts; and")
    add_bullet(doc, "referral or coupon credits applied to that shop.")

    add_subhead(doc, "9.6 Rules on deductions")
    add_bullet(
        doc,
        "Deductions must be real costs or the written allocation method in Schedule A, including the Maintenance Fee, not an arbitrary cut.",
    )
    add_bullet(
        doc,
        "The Company shall show the deductions, including the Maintenance Fee line, on the monthly statement.",
    )
    add_bullet(
        doc,
        "If a Resource Cost cannot fairly be tied to one shop, it is allocated equally across all active paying shops on the Platform for that month, then applied only to the Partner’s Referred Shops.",
    )
    add_bullet(
        doc,
        "Resource Costs for a shop (including Maintenance Fee) cannot exceed that shop’s Gross Collected Revenue in the same month. The Partner Share for that shop in that month would then be GHS 0, not a negative balance, unless a prior refund creates a true clawback under Clause 10.4.",
    )

    # 10
    add_heading_styled(doc, "10. Payment")
    add_para(doc, "10.1 Calculation (per Referred Shop, per month):")
    add_para(
        doc,
        "Partner Share = 50% × Net Distributable Revenue (after refunds, Paystack fees, Resource Costs, and Maintenance Fee).",
        italic=True,
        align="center",
    )
    add_para(
        doc,
        "10.2 The Company shall send a statement within 10 business days after month-end.",
    )
    add_para(
        doc,
        "10.3 The Company shall pay the Partner Share within 15 business days after the statement, by mobile money or bank transfer to:",
    )
    add_para(doc, "Account / MoMo name: Mr. John Paul Kpantey (or as he nominates in writing)", align="left")
    add_para(doc, "Number: ____________________", align="left")
    add_para(doc, "Bank: ____________________", align="left")
    add_para(
        doc,
        "10.4 If a later refund or chargeback relates to a period already paid, the Company may deduct it from the next Partner Share or invoice the Partner if none is due.",
    )
    add_para(
        doc,
        "10.5 If the monthly Partner Share is under GHS 50, the Company may roll it to the next month until the unpaid total is at least GHS 50, or until this Agreement ends, when any remainder is paid. Lifetime shop shares that continue after termination are still paid on this timetable.",
    )
    add_para(
        doc,
        "10.6 The Partner is responsible for his own income tax, GRA filings, and social security (if any). The Company may withhold tax if required by law and shall give evidence of withholding on request.",
    )

    # 11
    add_heading_styled(doc, "11. Worked example (illustrative only)")
    add_para(
        doc,
        "Assume one Referred Shop on Growth. The Company actually receives GHS 249.00.",
    )
    make_table(
        doc,
        ["Item", "Amount (GHS)"],
        [
            ["Gross Collected Revenue", "249.00"],
            ["Less Paystack fees (example 1.95%)", "(4.86)"],
            ["Less Maintenance Fee (Company updates & upkeep)", "(15.00)"],
            ["Less allocated hosting, database, Redis", "(8.00)"],
            ["Less allocated email, SMS, file storage", "(4.50)"],
            ["Less allocated monitoring, backups, domain share", "(2.00)"],
            ["Net Distributable Revenue", "214.64"],
            ["Partner Share (50%) — Mr. John Paul Kpantey", "107.32"],
            ["Company Share (50%) — ZeoFex Technology", "107.32"],
        ],
    )
    add_para(
        doc,
        "Figures in this example are illustrative. Live Paystack fees and Schedule A (including the agreed Maintenance Fee) control.",
        italic=True,
        size=11,
    )
    add_para(
        doc,
        "If the same shop is on Starter (GHS 0), Gross Collected Revenue is 0, so Partner Share is 0 until they pay. When they start paying, the lifetime 50% applies from that point. This 50% continues every month for the life of that shop, including after plan upgrades, as long as Clause 5.5 does not apply.",
    )

    # 12
    add_heading_styled(doc, "12. Intellectual property and brand")
    add_para(
        doc,
        "12.1 ZeoFex Technology owns VentraPOS, all software, updates, upgrades, trademarks, shop-platform data rights as under the VentraPOS Terms, and all related intellectual property.",
    )
    add_para(
        doc,
        "12.2 The Partner may use the VentraPOS name and approved sales materials only to originate shops, and must stop holding himself out as an active originator on termination, except as needed for an orderly handover. Lifetime payment on existing Referred Shops continues as in Clause 15.4.",
    )
    add_para(
        doc,
        "12.3 Shop operational data belongs to the shop and/or the Company under the VentraPOS Terms of Service. The Partner has no licence to copy, sell, or reuse that data.",
    )

    # 13
    add_heading_styled(doc, "13. Confidentiality")
    add_para(
        doc,
        "Each Party shall keep confidential the other Party’s non-public business information, including shop lists, pricing not on the public site, deduction workings, Maintenance Fee workings, and this Agreement’s commercial terms, for 3 years after termination, except where disclosure is required by law.",
    )

    # 14
    add_heading_styled(doc, "14. Non-solicit")
    add_para(
        doc,
        "During this Agreement and for 12 months after it ends, the Partner shall not knowingly induce a Referred Shop to leave VentraPOS for a competing POS product that the Partner owns, represents, or is paid to sell, except with the Company’s written consent. A breach of this clause ends future Partner Share on the affected shop(s) and may end remaining lifetime shares under Clause 15.4.",
    )

    # 15
    add_heading_styled(doc, "15. Term and termination")
    add_para(
        doc,
        "15.1 This Agreement starts on the Effective Date and continues until terminated.",
    )
    add_para(
        doc,
        "15.2 Either Party may terminate the active origination relationship for convenience on 30 days’ written notice. Termination of origination does not by itself cancel lifetime Partner Share on shops already recorded as Referred Shops.",
    )
    add_para(
        doc,
        "15.3 Either Party may terminate immediately if the other materially breaches and does not remedy within 14 days of written notice, or if the other becomes insolvent, or if the Partner commits fraud, bribery, or misuse of shop data.",
    )
    add_para(doc, "15.4 On termination:")
    add_bullet(
        doc,
        "the Partner shall stop representing himself as an active VentraPOS originating partner for new shops;",
    )
    add_bullet(doc, "the Company shall pay any undisputed Partner Share then due;")
    add_bullet(
        doc,
        "existing Referred Shops continue to generate Partner Share for the Life of each Shop (Clause 5), including Maintenance Fee deductions as usual;",
    )
    add_bullet(
        doc,
        "exception: if termination is for the Partner’s fraud, data misuse, or Clause 14 breach, future Partner Share stops.",
    )
    add_para(
        doc,
        "15.5 After termination of origination, the Company remains solely responsible for all updates and maintenance. The Maintenance Fee continues to be deducted as a Resource Cost for as long as Partner Share is paid on that shop.",
    )

    # 16
    add_heading_styled(doc, "16. Liability")
    add_para(
        doc,
        "16.1 Neither Party is liable to the other for indirect or consequential loss, including lost introductions that did not convert.",
    )
    add_para(
        doc,
        "16.2 The Company’s total liability under this Agreement in any 12-month period is limited to the Partner Share actually paid or payable in that period.",
    )
    add_para(
        doc,
        "16.3 Nothing limits liability for fraud or death or personal injury caused by negligence.",
    )

    # 17
    add_heading_styled(doc, "17. General")
    add_para(
        doc,
        "17.1 This is the entire agreement on this subject. Changes must be in writing and signed by both Parties.",
    )
    add_para(
        doc,
        "17.2 The Partner may not assign this Agreement without the Company’s written consent, except the estate provision in Clause 5.6. The Company may assign to a successor of the VentraPOS business. A sale of the business does not cancel lifetime Partner Share on existing Referred Shops; the successor must honour Clause 5.",
    )
    add_para(
        doc,
        "17.3 Notices may be sent to the emails above and are deemed received the next business day.",
    )
    add_para(doc, "17.4 If a clause is invalid, the rest remains in force.")
    add_para(
        doc,
        "17.5 The laws of Ghana govern this Agreement. The courts of Ghana have exclusive jurisdiction. The Parties may instead agree in writing to mediation in Accra, then arbitration under the Alternative Dispute Resolution Act, 2010 (Act 798).",
    )

    # Schedule A
    add_heading_styled(doc, "Schedule A — Resource Cost Allocation (including Maintenance Fee)")
    add_para(doc, "For each Settlement Period, Resource Costs for a Referred Shop equal:")
    add_para(
        doc,
        "(1) Direct costs tagged to that shop (SMS OTPs for that shop’s users, extra storage, dedicated support if sold); plus (2) Maintenance Fee for that shop; plus (3) Shared Platform costs for that month divided by the number of active paying shops on VentraPOS that month.",
    )
    add_para(doc, "Maintenance Fee (agreed method) — choose and initial one:", bold=True)
    add_para(
        doc,
        "☐  Fixed: GHS ________ per paying Referred Shop per month;    or",
        align="left",
    )
    add_para(
        doc,
        "☐  Percentage: ________% of that shop’s Gross Collected Revenue;    or",
        align="left",
    )
    add_para(
        doc,
        "☐  Allocated actuals: Company engineering / maintenance cost for the month ÷ active paying shops.",
        align="left",
    )
    add_para(
        doc,
        "Until a box is ticked, the Company shall use allocated actuals and show the figure on the monthly statement.",
        italic=True,
        size=11,
    )
    add_para(doc, "Shared Platform cost categories (fill actuals each month or keep a running method):")
    make_table(
        doc,
        ["Category", "Vendor / notes", "Monthly actual (GHS)"],
        [
            ["Maintenance Fee / updates (if not a per-shop fixed fee)", "ZeoFex Technology — software updates & upkeep", ""],
            ["Cloud hosting & compute", "", ""],
            ["Database", "PostgreSQL hosting & backups", ""],
            ["Redis / queues", "", ""],
            ["Email", "", ""],
            ["SMS", "", ""],
            ["File storage", "", ""],
            ["Domain / SSL / DNS", "", ""],
            ["Monitoring", "", ""],
            ["Payment processing", "Paystack (usually deducted per transaction under 9.2, not here)", "—"],
            ["Total shared (excluding per-transaction Paystack)", "", ""],
            ["÷ active paying shops", "", ""],
            ["Allocated Resource Cost per paying shop", "", ""],
        ],
    )
    add_para(
        doc,
        "Founder drawings, general marketing, and brand-new product lines outside ordinary VentraPOS maintenance are not Resource Costs unless the Parties add them in writing.",
        italic=True,
        size=11,
    )

    # Schedule B
    add_heading_styled(doc, "Schedule B — Commercial terms")
    make_table(
        doc,
        ["Item", "Agreed term"],
        [
            ["Company", "ZeoFex Technology (VentraPOS)"],
            ["Partner", "Mr. John Paul Kpantey"],
            ["Partner Share", "50% of Net Distributable Revenue per Referred Shop"],
            ["Duration", "For the life of each Referred Shop on VentraPOS"],
            ["Updates", "Company only. Partner does not build or ship software"],
            ["Maintenance Fee", "Deducted as a Resource Cost before the 50/50 split"],
            ["Does not apply to", "Free Starter shops until they pay; shops the Partner did not bring"],
            ["Settlement", "Monthly; pay within 15 business days of statement"],
            ["Currency", "Ghana Cedis (GHS)"],
            ["Tracking method", "________________________________"],
            ["Authorised discount ceiling", "________________________________"],
            ["Hardware / scanners / printers", "Excluded unless added in writing"],
        ],
    )

    # Signatures
    add_heading_styled(doc, "Signature")
    add_para(doc, "Signed as a binding agreement.", space_after=14)

    add_para(doc, "FOR ZEOFEX TECHNOLOGY (trading as VentraPOS)", bold=True, align="left", space_after=12)
    add_blank_line(doc, "Name:")
    add_blank_line(doc, "Title:")
    add_blank_line(doc, "Signature:")
    add_blank_line(doc, "Date:")
    add_blank_line(doc, "Witness name / signature:")

    add_para(doc, "FOR THE PARTNER", bold=True, align="left", space_before=16, space_after=12)
    add_para(doc, "Name: Mr. John Paul Kpantey", align="left", space_after=10)
    add_blank_line(doc, "Signature:")
    add_blank_line(doc, "Date:")
    add_blank_line(doc, "Ghana Card / ID:")
    add_blank_line(doc, "Witness name / signature:")

    add_para(
        doc,
        "This document is a commercial draft prepared for the Parties. It should be reviewed by a lawyer licensed in Ghana before it is signed.",
        italic=True,
        size=10,
        color=GRAY,
        space_before=18,
        align="center",
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
